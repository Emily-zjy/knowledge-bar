from __future__ import annotations

import base64
import binascii
import logging
import re
from io import BytesIO
from urllib.parse import quote

from docx import Document
from docx.shared import Inches
from fastapi import APIRouter, Header
from fastapi.responses import JSONResponse, PlainTextResponse, StreamingResponse

from app.core.config import settings
from app.models.ai import (
    ApiError,
    ApiResponse,
    ExportWordRequest,
    GenerateMindMapRequest,
    GenerateMindMapResult,
    GenerationJobResponse,
    LatestMindMapResponse,
    MindMapData,
    MindMapVersionItem,
    MindMapVersionsResponse,
    PatchMindMapRequest,
    RegenerateMindMapRequest,
)
from app.services.ai_mindmap import MindMapGenerationError, generate_mind_map
from app.services.auth import AuthError, verify_required_supabase_jwt
from app.services.supabase_store import SupabaseStoreError, supabase_store

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/ai", tags=["ai"])
IMAGE_MARKDOWN_PATTERN = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)(?:\{width=(\d+)\})?")


def _api_error(status_code: int, code: str, message: str) -> JSONResponse:
    return JSONResponse(
        status_code=status_code,
        content=ApiResponse(success=False, error=ApiError(code=code, message=message)).model_dump(),
    )


def _handle_auth_error(exc: AuthError) -> JSONResponse:
    status = 503 if exc.code in ("AUTH_NOT_CONFIGURED", "AUTH_UNAVAILABLE") else 401
    return _api_error(status, exc.code, exc.message)


def _public_generation_error_message(exc: MindMapGenerationError) -> str:
    if exc.error_stage in {"provider_5xx", "provider_timeout", "provider_network_error"}:
        return "AI 服务暂时繁忙，请稍后再试。"
    if exc.error_stage in {"json_parse", "schema_validation", "schema_repair_failed"}:
        return "AI 返回结果异常，请重新生成。"
    if exc.error_stage == "provider_4xx":
        return "AI 服务请求失败，请稍后再试。"
    return "思维导图生成失败，请重新生成。"


def mind_map_to_markdown(mind_map: MindMapData) -> str:
    children_by_parent: dict[str | None, list] = {}
    for node in mind_map.nodes:
        children_by_parent.setdefault(node.parentId, []).append(node)
    for siblings in children_by_parent.values():
        siblings.sort(key=lambda node: node.order)

    lines: list[str] = []

    def render(parent_id: str | None, level: int) -> None:
        for node in children_by_parent.get(parent_id, []):
            if level == 0:
                lines.append(f"# {node.text}")
            else:
                lines.append(f"{'  ' * max(level - 1, 0)}- {node.text}")
            render(node.id, level + 1)

    render(None, 0)
    return "\n".join(lines)


def _add_markdown_image(document: Document, alt: str, src: str, width: str | None = None) -> bool:
    if not src.startswith("data:image/"):
        return False

    try:
        _, encoded = src.split(",", 1)
        image_bytes = base64.b64decode(encoded)
    except (ValueError, binascii.Error):
        return False

    width_inches = None
    if width:
        try:
            width_pixels = min(max(int(width), 80), 640)
            width_inches = Inches(width_pixels / 96)
        except ValueError:
            width_inches = None

    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    run.add_picture(BytesIO(image_bytes), width=width_inches)
    if alt:
        document.add_paragraph(alt)
    return True


def _add_markdown_line(document: Document, line: str) -> None:
    remaining = line
    while True:
        match = IMAGE_MARKDOWN_PATTERN.search(remaining)
        if not match:
            break

        prefix = remaining[: match.start()].strip()
        if prefix:
            document.add_paragraph(prefix)

        alt, src, width = match.group(1), match.group(2), match.group(3)
        if not _add_markdown_image(document, alt, src, width):
            document.add_paragraph(alt or "image")

        remaining = remaining[match.end() :]

    if remaining.strip():
        document.add_paragraph(remaining.strip())


def _word_stream(title: str, markdown: str) -> StreamingResponse:
    document = Document()
    document.add_heading(title, level=1)

    if markdown.strip():
        for raw_line in markdown.splitlines():
            line = raw_line.rstrip()
            if not line:
                document.add_paragraph("")
            elif line.startswith("# "):
                document.add_heading(line[2:].strip(), level=1)
            elif line.startswith("## "):
                document.add_heading(line[3:].strip(), level=2)
            elif line.lstrip().startswith("- "):
                document.add_paragraph(line.lstrip()[2:].strip(), style="List Bullet")
            else:
                _add_markdown_line(document, line)
    else:
        document.add_paragraph("No content.")

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    filename = quote(f"{title}.docx")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


# === Mind Map Generation ===


@router.post("/mindmap/generate", response_model=None)
def generate_mindmap(
    payload: GenerateMindMapRequest,
    authorization: str | None = Header(default=None),
) -> ApiResponse[GenerateMindMapResult] | JSONResponse:
    job_id: str | None = None
    user_id: str | None = None
    try:
        user_id = verify_required_supabase_jwt(authorization)
        kb = supabase_store.get_knowledge_base_for_user(str(payload.knowledge_base_id), user_id)
        content_md = payload.content_md.strip() or str(kb.get("content_md") or "").strip()
        if not content_md:
            return _api_error(422, "EMPTY_CONTENT", "Knowledge content cannot be empty.")
        if len(content_md) > settings.max_content_length:
            return _api_error(422, "CONTENT_TOO_LONG", "Content exceeds the maximum allowed length.")

        db_content = str(kb.get("content_md") or "").strip()
        if not db_content:
            supabase_store.update_knowledge_base_content(str(payload.knowledge_base_id), user_id, content_md)

        job = supabase_store.create_generation_job(user_id, str(payload.knowledge_base_id), "running")
        job_id = str(job["id"])
        generation_payload = payload.model_copy(
            update={
                "title": str(kb.get("name") or payload.title),
                "description": str(kb.get("description") or payload.description or ""),
                "content_md": content_md,
            },
        )
        result = generate_mind_map(payload=generation_payload, authorization=authorization, guest_session_id=None)
        result.generation_id = job_id
        markdown = mind_map_to_markdown(result.mind_map)
        supabase_store.save_mind_map(
            user_id=user_id,
            knowledge_base_id=str(payload.knowledge_base_id),
            mind_map=result.mind_map,
            markdown=markdown,
        )
        meta = result.metadata
        supabase_store.update_generation_job(
            job_id, user_id, status="succeeded",
            input_length=meta.input_length if meta else None,
            input_tokens=meta.input_tokens if meta else None,
            output_tokens=meta.output_tokens if meta else None,
            retry_count=meta.retry_count if meta else None,
            schema_repair_count=meta.schema_repair_count if meta else None,
            repair_attempted=meta.repair_attempted if meta else None,
            repair_success=meta.repair_success if meta else None,
            latency_ms=meta.latency_ms if meta else None,
            node_count=meta.node_count if meta else None,
            tree_depth=meta.tree_depth if meta else None,
            requested_depth=payload.requested_max_depth,
            prompt_version=meta.prompt_version if meta else None,
        )
        result.metadata = None
        return ApiResponse(success=True, data=result)
    except AuthError as exc:
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        return _api_error(exc.status_code, exc.code, exc.message)
    except MindMapGenerationError as exc:
        if job_id and user_id:
            try:
                supabase_store.update_generation_job(
                    job_id,
                    user_id,
                    status="failed",
                    error_code=exc.code,
                    error_message=exc.message,
                    error_stage=exc.error_stage,
                    latency_ms=exc.latency_ms,
                    retry_count=exc.retry_count,
                    input_length=exc.input_length,
                    requested_depth=payload.requested_max_depth,
                    repair_attempted=exc.repair_attempted,
                    repair_success=exc.repair_success,
                    prompt_version=exc.prompt_version,
                )
            except SupabaseStoreError:
                logger.exception("failed_to_update_generation_job")
        logger.warning("mindmap_generate_failed", extra={
            "error_stage": exc.error_stage, "code": exc.code,
            "latency_ms": exc.latency_ms, "retry_count": exc.retry_count,
        })
        return _api_error(502, exc.code, _public_generation_error_message(exc))
# === Regenerate ===


@router.post("/mindmap/{knowledge_base_id}/regenerate", response_model=None)
def regenerate_mindmap(
    knowledge_base_id: str,
    body: RegenerateMindMapRequest | None = None,
    authorization: str | None = Header(default=None),
) -> ApiResponse[GenerateMindMapResult] | JSONResponse:
    job_id: str | None = None
    user_id: str | None = None
    try:
        user_id = verify_required_supabase_jwt(authorization)
        kb = supabase_store.get_knowledge_base_for_user(knowledge_base_id, user_id)
        content_md = str(kb.get("content_md") or "").strip()
        if not content_md and body and body.content_md:
            content_md = body.content_md.strip()
            if content_md:
                supabase_store.update_knowledge_base_content(knowledge_base_id, user_id, content_md)
        if not content_md:
            logger.warning(
                "mindmap_regenerate_failed",
                extra={"error_stage": "request_validation", "reason": "empty_content_md",
                       "knowledge_base_id": knowledge_base_id, "status_code": 422},
            )
            return _api_error(422, "EMPTY_CONTENT", "Knowledge content cannot be empty. Please save your content before regenerating.")

        requested_depth = body.requested_max_depth if body else 5
        job = supabase_store.create_generation_job(user_id, knowledge_base_id, "running")
        job_id = str(job["id"])
        generation_payload = GenerateMindMapRequest(
            knowledge_base_id=knowledge_base_id,
            title=str(kb.get("name") or "Untitled"),
            description=str(kb.get("description") or ""),
            content_md=content_md,
            requested_max_depth=requested_depth,
        )
        result = generate_mind_map(payload=generation_payload, authorization=authorization, guest_session_id=None)
        result.generation_id = job_id
        markdown = mind_map_to_markdown(result.mind_map)
        supabase_store.save_mind_map(
            user_id=user_id,
            knowledge_base_id=knowledge_base_id,
            mind_map=result.mind_map,
            markdown=markdown,
        )
        meta = result.metadata
        supabase_store.update_generation_job(
            job_id, user_id, status="succeeded",
            input_length=meta.input_length if meta else None,
            input_tokens=meta.input_tokens if meta else None,
            output_tokens=meta.output_tokens if meta else None,
            retry_count=meta.retry_count if meta else None,
            schema_repair_count=meta.schema_repair_count if meta else None,
            repair_attempted=meta.repair_attempted if meta else None,
            repair_success=meta.repair_success if meta else None,
            latency_ms=meta.latency_ms if meta else None,
            node_count=meta.node_count if meta else None,
            tree_depth=meta.tree_depth if meta else None,
            requested_depth=requested_depth,
            prompt_version=meta.prompt_version if meta else None,
        )
        logger.info("mindmap_regenerate_succeeded", extra={"knowledge_base_id": knowledge_base_id, "job_id": job_id})
        result.metadata = None
        return ApiResponse(success=True, data=result)
    except AuthError as exc:
        logger.warning("mindmap_regenerate_failed", extra={"error_stage": "auth", "knowledge_base_id": knowledge_base_id})
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        logger.warning(
            "mindmap_regenerate_failed",
            extra={"error_stage": "knowledge_base", "knowledge_base_id": knowledge_base_id,
                   "status_code": exc.status_code, "code": exc.code},
        )
        return _api_error(exc.status_code, exc.code, exc.message)
    except MindMapGenerationError as exc:
        logger.warning(
            "mindmap_regenerate_failed",
            extra={"error_stage": exc.error_stage, "knowledge_base_id": knowledge_base_id,
                   "code": exc.code, "latency_ms": exc.latency_ms},
        )
        if job_id and user_id:
            try:
                supabase_store.update_generation_job(
                    job_id,
                    user_id,
                    status="failed",
                    error_code=exc.code,
                    error_message=exc.message,
                    error_stage=exc.error_stage,
                    latency_ms=exc.latency_ms,
                    retry_count=exc.retry_count,
                    input_length=exc.input_length,
                    requested_depth=body.requested_max_depth if body else 5,
                    repair_attempted=exc.repair_attempted,
                    repair_success=exc.repair_success,
                    prompt_version=exc.prompt_version,
                )
            except SupabaseStoreError:
                logger.exception("failed_to_update_generation_job")
        return _api_error(502, exc.code, _public_generation_error_message(exc))


# === Job Status ===


@router.get("/mindmap/jobs/{job_id}", response_model=None)
def get_generation_job(
    job_id: str,
    authorization: str | None = Header(default=None),
) -> ApiResponse[GenerationJobResponse] | JSONResponse:
    try:
        user_id = verify_required_supabase_jwt(authorization)
        row = supabase_store.get_generation_job_for_user(job_id, user_id)
        return ApiResponse(success=True, data=GenerationJobResponse(**row))
    except AuthError as exc:
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        return _api_error(exc.status_code, exc.code, exc.message)


# === Get Latest Mind Map ===


@router.get("/mindmap/latest/{knowledge_base_id}", response_model=None)
def get_latest_mind_map(
    knowledge_base_id: str,
    authorization: str | None = Header(default=None),
) -> ApiResponse[LatestMindMapResponse] | JSONResponse:
    try:
        user_id = verify_required_supabase_jwt(authorization)
        supabase_store.get_knowledge_base_for_user(knowledge_base_id, user_id)
        row = supabase_store.get_latest_mind_map(knowledge_base_id, user_id)
        data = MindMapData.model_validate(row["data_jsonb"])
        return ApiResponse(
            success=True,
            data=LatestMindMapResponse(
                id=str(row.get("id") or ""),
                knowledge_base_id=knowledge_base_id,
                version_no=int(row.get("version_no") or 1),
                mind_map=data,
                markdown=str(row.get("markdown") or ""),
                updated_at=row.get("updated_at"),
            ),
        )
    except AuthError as exc:
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        return _api_error(exc.status_code, exc.code, exc.message)


# === Version List ===


@router.get("/mindmap/versions/{knowledge_base_id}", response_model=None)
def list_mind_map_versions(
    knowledge_base_id: str,
    authorization: str | None = Header(default=None),
) -> ApiResponse[MindMapVersionsResponse] | JSONResponse:
    try:
        user_id = verify_required_supabase_jwt(authorization)
        supabase_store.get_knowledge_base_for_user(knowledge_base_id, user_id)
        rows = supabase_store.list_mind_map_versions(knowledge_base_id, user_id)
        versions = [
            MindMapVersionItem(
                id=str(r.get("id") or ""),
                version_no=int(r.get("version_no") or 1),
                updated_at=r.get("updated_at"),
            )
            for r in rows
        ]
        return ApiResponse(success=True, data=MindMapVersionsResponse(knowledge_base_id=knowledge_base_id, versions=versions))
    except AuthError as exc:
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        return _api_error(exc.status_code, exc.code, exc.message)


# === Patch Mind Map (manual edit) ===


@router.patch("/mindmap/{mind_map_id}", response_model=None)
def patch_mind_map(
    mind_map_id: str,
    payload: PatchMindMapRequest,
    authorization: str | None = Header(default=None),
) -> ApiResponse[LatestMindMapResponse] | JSONResponse:
    try:
        user_id = verify_required_supabase_jwt(authorization)
        existing = supabase_store.get_mind_map_by_id(mind_map_id, user_id)
        kb_id = str(existing.get("knowledge_base_id") or "")
        supabase_store.get_knowledge_base_for_user(kb_id, user_id)

        current_data = MindMapData.model_validate(existing["data_jsonb"])
        nodes = current_data.nodes

        if payload.update_node:
            node_id = payload.update_node.id
            new_text = payload.update_node.text
            for node in nodes:
                if node.id == node_id:
                    node.text = new_text[:120]
                    break

        if payload.title:
            current_data.title = payload.title

        updated_data = MindMapData(title=current_data.title, nodes=nodes)
        markdown = mind_map_to_markdown(updated_data)
        row = supabase_store.patch_mind_map(mind_map_id, user_id, updated_data.model_dump(), markdown)

        return ApiResponse(
            success=True,
            data=LatestMindMapResponse(
                id=mind_map_id,
                knowledge_base_id=kb_id,
                version_no=int(row.get("version_no") or 1),
                mind_map=updated_data,
                markdown=markdown,
                updated_at=row.get("updated_at"),
            ),
        )
    except AuthError as exc:
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        return _api_error(exc.status_code, exc.code, exc.message)


# === Delete Mind Map ===


@router.delete("/mindmap/{mind_map_id}", response_model=None)
def delete_mind_map(
    mind_map_id: str,
    authorization: str | None = Header(default=None),
) -> ApiResponse | JSONResponse:
    try:
        user_id = verify_required_supabase_jwt(authorization)
        existing = supabase_store.get_mind_map_by_id(mind_map_id, user_id)
        kb_id = str(existing.get("knowledge_base_id") or "")
        supabase_store.get_knowledge_base_for_user(kb_id, user_id)
        supabase_store.delete_mind_map(mind_map_id, user_id)
        return ApiResponse(success=True)
    except AuthError as exc:
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        return _api_error(exc.status_code, exc.code, exc.message)


# === Export: Markdown ===


@router.get("/export/markdown/{knowledge_base_id}", response_model=None)
def export_markdown(
    knowledge_base_id: str,
    authorization: str | None = Header(default=None),
) -> PlainTextResponse | JSONResponse:
    try:
        user_id = verify_required_supabase_jwt(authorization)
        kb = supabase_store.get_knowledge_base_for_user(knowledge_base_id, user_id)
        try:
            latest = supabase_store.get_latest_mind_map(knowledge_base_id, user_id)
            markdown = str(latest.get("markdown") or "") or mind_map_to_markdown(MindMapData.model_validate(latest["data_jsonb"]))
        except SupabaseStoreError as exc:
            if exc.status_code != 404:
                raise
            markdown = str(kb.get("content_md") or "")
        filename = quote(f"{kb.get('name') or 'knowledge-bar'}.md")
        return PlainTextResponse(
            markdown,
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
        )
    except AuthError as exc:
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        return _api_error(exc.status_code, exc.code, exc.message)


# === Export: Word (GET with KB ID) ===


@router.get("/export/word/{knowledge_base_id}", response_model=None)
def export_word_from_knowledge_base(
    knowledge_base_id: str,
    authorization: str | None = Header(default=None),
) -> StreamingResponse | JSONResponse:
    try:
        user_id = verify_required_supabase_jwt(authorization)
        kb = supabase_store.get_knowledge_base_for_user(knowledge_base_id, user_id)
        try:
            latest = supabase_store.get_latest_mind_map(knowledge_base_id, user_id)
            markdown = str(latest.get("markdown") or "") or mind_map_to_markdown(MindMapData.model_validate(latest["data_jsonb"]))
        except SupabaseStoreError as exc:
            if exc.status_code != 404:
                raise
            markdown = str(kb.get("content_md") or "")
        return _word_stream(str(kb.get("name") or "Knowledge Bar"), markdown)
    except AuthError as exc:
        return _handle_auth_error(exc)
    except SupabaseStoreError as exc:
        return _api_error(exc.status_code, exc.code, exc.message)


# === Export: Word (POST with inline markdown) ===


@router.post("/export/word", response_model=None)
def export_word(
    payload: ExportWordRequest,
    authorization: str | None = Header(default=None),
) -> StreamingResponse | JSONResponse:
    if payload.knowledge_base_id:
        try:
            user_id = verify_required_supabase_jwt(authorization)
            supabase_store.get_knowledge_base_for_user(str(payload.knowledge_base_id), user_id)
        except AuthError as exc:
            return _handle_auth_error(exc)
        except SupabaseStoreError as exc:
            return _api_error(exc.status_code, exc.code, exc.message)

    if len(payload.markdown) > settings.max_export_length:
        return _api_error(422, "CONTENT_TOO_LONG", "Export content exceeds the maximum allowed length.")

    return _word_stream(payload.title, payload.markdown)
